"""Turn an uploaded resume file into plain text. The file itself is never stored."""

from __future__ import annotations

import io
import re
import zipfile

MAX_BYTES = 5 * 1024 * 1024
MAX_PAGES = 10


class UnreadableResume(ValueError):
    """The file isn't a resume we can read; the message is safe to show."""


def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("", "•").replace("•", "•").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------- links ----------
#
# Resumes often show "GitHub • LinkedIn • Portfolio" with the address hidden
# behind the words. Text extraction keeps only the words, so the addresses
# are read from the file's links and written into the text next to them:
# then the contact check and the builder see the same links a reader would.

HEADER_SHARE = 0.25  # links in the top quarter of the first page are the contact links
_LABEL_LINE = re.compile(r"linkedin|github|portfolio|website|blog|gitlab|behance|dribbble", re.I)


def _web_links(uris: list[str], text: str) -> list[str]:
    """http(s) addresses not already written in the text, in order, once each."""
    seen = []
    for uri in uris:
        uri = uri.strip()
        bare = re.sub(r"^https?://(www\.)?", "", uri).rstrip("/")
        if not re.match(r"^https?://", uri, re.I) or not bare or bare.lower() in text.lower() or uri in seen:
            continue
        seen.append(uri)
    return seen


def _with_header_links(text: str, links: list[str]) -> str:
    """Put header links on their own line, right after the line that names them
    ("GitHub • LinkedIn") or, if no line does, after the name."""
    if not links:
        return text
    lines = text.split("\n")
    head = min(len(lines), 8)
    at = next((i for i in range(head) if _LABEL_LINE.search(lines[i])), 0)
    lines.insert(at + 1, " | ".join(links))
    return "\n".join(lines)


def _pdf_header_links(page) -> list[str]:
    """Link addresses in the top part of a PDF page."""
    try:
        top = float(page.mediabox.top)
        height = float(page.mediabox.height) or 1.0
        annotations = page.get("/Annots") or []
        found = []
        for ref in annotations:
            annot = ref.get_object()
            action = annot.get("/A") or {}
            uri = action.get("/URI") if hasattr(action, "get") else None
            rect = annot.get("/Rect")
            if not uri or not rect:
                continue
            if (top - float(max(rect[1], rect[3]))) / height <= HEADER_SHARE:
                found.append(str(uri))
        return found
    except (KeyError, TypeError, ValueError, AttributeError):
        return []  # a malformed annotation never stops the text from being read


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadableResume(
                "This PDF is password protected. Save a copy without a password and try again."
            )
        pages = reader.pages[:MAX_PAGES]
        text = "\n\n".join(page.extract_text() or "" for page in pages)
        links = _pdf_header_links(pages[0]) if pages else []
        return _with_header_links(text, _web_links(links, text))
    except PdfReadError as exc:
        raise UnreadableResume(
            "This PDF couldn't be read. Try exporting it again, or upload a Word file."
        ) from exc


def _docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise UnreadableResume("This Word file couldn't be read. Save it again as .docx and retry.") from exc
    lines = []
    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        prefix = "• " if "list" in style else ""
        # a hyperlink's text is in paragraph.text, its address isn't: add it on the same line
        addresses = [link.address for link in paragraph.hyperlinks if link.address]
        extra = _web_links(addresses, paragraph.text)
        lines.append(prefix + paragraph.text + (" | " + " | ".join(extra) if extra else ""))
    for table in document.tables:  # some templates lay resumes out in tables
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def extract_text(data: bytes, filename: str = "", content_type: str = "") -> str:
    if not data:
        raise UnreadableResume("The file is empty.")
    if len(data) > MAX_BYTES:
        raise UnreadableResume("Resumes can be up to 5 MB.")
    name = filename.lower()
    if data[:4] == b"%PDF" or name.endswith(".pdf"):
        text = _pdf(data)
    elif data[:2] == b"PK" or name.endswith(".docx"):
        text = _docx(data)
    elif name.endswith((".txt", ".md")) or content_type.startswith("text/"):
        text = data.decode("utf-8", errors="replace")
    elif name.endswith(".doc"):
        raise UnreadableResume("Old .doc files aren't supported. Save it as .docx or PDF and try again.")
    else:
        raise UnreadableResume("Upload a PDF or Word (.docx) file.")
    text = _clean(text)
    if len(text) < 200:
        raise UnreadableResume(
            "Almost no text came out of this file. If it's a scanned image, export a text PDF from your editor instead."
        )
    return text[:60_000]
