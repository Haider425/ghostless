"""The resume engine: reading files, building a profile, checks and scoring."""

import io
import os
from datetime import date

import pytest

from tracker.resume.document import ResumeDocument, document_to_text
from tracker.resume.extract import UnreadableResume, extract_text
from tracker.resume.match import recommend, score_job, title_words
from tracker.resume.profile import build_profile, date_ranges, months_of_experience, split_sections
from tracker.resume.quality import check_resume

TODAY = date(2026, 10, 17)

STRONG = """Jordan Lee
jordan.lee@example.com | 416-555-0199 | Toronto, ON
linkedin.com/in/jordanlee | github.com/jordanlee

EDUCATION
University of Waterloo | Bachelor of Computer Science | Expected April 2027

EXPERIENCE
Software Developer Co-op | Maple Bank | May 2025 - Aug 2025
• Built a React and TypeScript dashboard used by 40 analysts, cutting report time by 30%
• Automated nightly SQL Server jobs with Python, saving 6 hours a week
• Designed REST APIs in C# and ASP.NET for 3 internal tools

Data Engineering Intern | Northwind | Jan 2026 - Apr 2026
• Migrated 12 Airflow pipelines to Spark on Databricks, reducing runtime by 45%
• Wrote unit tests with pytest covering 85% of the ingestion code

PROJECTS
Job tracker | Python, PostgreSQL, Docker, AWS
• Scraped 50,000 postings a day and served them through a FastAPI service

SKILLS
Languages: Python, TypeScript, C#, SQL, Java
Tools: React, Docker, AWS, Git, Linux, Kubernetes
"""

WEAK = """Alex Kim
Toronto

Experience
Retail Associate | Shop | 2023 - 2024
• Responsible for helping customers
• Helped with inventory and other tasks around the store
• Worked on the cash register during busy times

I am a hardworking team player and passionate about technology.
"""


# ---------- files ----------


def _pdf(lines: list[str], links: tuple[tuple[str, int], ...] = ()) -> bytes:
    """A minimal one-page PDF with text, built by hand. `links` are
    (address, height on the page) pairs, added as clickable link annotations."""
    content = (
        "BT /F1 11 Tf 50 750 Td 14 TL "
        + " ".join(f"({line.replace('(', '').replace(')', '')}) Tj T*" for line in lines)
        + " ET"
    )
    annots = " ".join(f"{6 + i} 0 R" for i in range(len(links)))
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        f"/Resources << /Font << /F1 5 0 R >> >> /Annots [{annots}] >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        *(
            f"<< /Type /Annot /Subtype /Link /Rect [50 {y} 120 {y + 12}] /A << /S /URI /URI ({uri}) >> >>"
            for uri, y in links
        ),
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def test_reads_pdf():
    text = extract_text(_pdf([line.replace("•", "-") for line in STRONG.splitlines()[:20]]), "resume.pdf")
    assert "Jordan Lee" in text and "Databricks" in text


LABELLED_HEADER = [
    "Haider Saleem",
    "haider@example.com - 647-555-0140 - GitHub - LinkedIn - Portfolio",
    "EDUCATION",
    "Ontario Tech University | Bachelor of Science in Computer Science | Graduating December 2027",
    "EXPERIENCE",
    "Software Developer Co-op | Ontario Power Generation | September 2026 - April 2027",
    "- Developed a scheduling web app with Angular and ASP.NET Core used by 20+ staff",
    "- Combined data from 4 databases with SQL to narrow 1 million records in under 4 seconds",
    "PROJECTS",
    "Job tracker | Python, PostgreSQL",
    "- Scraped 50,000 postings a day and served them through a FastAPI service",
]


def test_pdf_header_links_are_read_from_the_links_behind_the_words():
    pdf = _pdf(
        LABELLED_HEADER,
        links=(
            ("https://github.com/haider", 734),
            ("https://www.linkedin.com/in/haider", 734),
            ("https://haider.dev", 734),
            ("mailto:haider@example.com", 734),  # an email link is not a web link
            ("https://github.com/haider/job-tracker", 150),  # a project link, far below the header
        ),
    )
    lines = extract_text(pdf, "resume.pdf").splitlines()
    # right after the line that names them, so the builder reads them as contact links
    assert lines[1].endswith("GitHub - LinkedIn - Portfolio")
    assert lines[2] == "https://github.com/haider | https://www.linkedin.com/in/haider | https://haider.dev"
    assert not any("job-tracker" in line for line in lines)


def test_the_builder_gets_the_header_links_of_an_uploaded_pdf():
    from tracker.resume.to_document import text_to_document

    links = (
        ("https://github.com/haider", 734),
        ("https://www.linkedin.com/in/haider", 734),
        ("https://haider.dev", 734),
    )
    contact = text_to_document(extract_text(_pdf(LABELLED_HEADER, links=links), "resume.pdf")).contact
    assert [(link.label, link.url) for link in contact.links] == [
        ("GitHub", "https://github.com/haider"),
        ("LinkedIn", "https://www.linkedin.com/in/haider"),
        ("Website", "https://haider.dev"),
    ]


def test_links_with_their_address_already_written_are_not_repeated():
    lines = ["Jordan Lee", "jordan@example.com | linkedin.com/in/jordanlee", *LABELLED_HEADER[2:]]
    text = extract_text(_pdf(lines, links=(("https://linkedin.com/in/jordanlee", 734),)), "resume.pdf")
    assert text.count("linkedin.com/in/jordanlee") == 1


def test_a_resume_with_link_words_in_its_header_is_not_asked_for_links():
    text = extract_text(
        _pdf(
            LABELLED_HEADER,
            links=(("https://github.com/haider", 734), ("https://linkedin.com/in/haider", 734)),
        ),
        "r.pdf",
    )
    profile = build_profile(text, TODAY)
    assert profile.contact["linkedin"] and profile.contact["github"]
    checks, _ = check_resume(text, profile)
    codes = {check.id for check in checks}
    assert "no_linkedin" not in codes and "no_portfolio" not in codes
    # even pasted text that lost the addresses: the header still names them
    pasted = "\n".join(LABELLED_HEADER)
    assert build_profile(pasted, TODAY).contact["linkedin"]


def test_link_words_outside_the_header_do_not_count():
    text = "\n".join(
        LABELLED_HEADER[:1]
        + ["haider@example.com"]
        + LABELLED_HEADER[2:]
        + ["SKILLS", "GitHub Actions, LinkedIn Learning"]
    )
    contact = build_profile(text, TODAY).contact
    assert not contact["linkedin"] and not contact["github"]


def test_docx_hyperlink_addresses_are_kept():
    docx = pytest.importorskip("docx")
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    document = docx.Document()
    document.add_paragraph("Haider Saleem")
    contact = document.add_paragraph("haider@example.com | ")
    rel = contact.part.relate_to(
        "https://www.linkedin.com/in/haider", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    words = OxmlElement("w:t")
    words.text = "LinkedIn"
    run.append(words)
    link.append(run)
    contact._p.append(link)
    for line in LABELLED_HEADER[2:]:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    lines = extract_text(buffer.getvalue(), "resume.docx").splitlines()
    assert lines[1] == "haider@example.com | LinkedIn | https://www.linkedin.com/in/haider"


def test_reads_docx():
    docx = pytest.importorskip("docx")
    document = docx.Document()
    for line in STRONG.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    assert "Maple Bank" in extract_text(buffer.getvalue(), "resume.docx")


@pytest.mark.parametrize(
    "data, name, message",
    [
        (b"", "r.pdf", "empty"),
        (b"hello", "r.exe", "PDF or Word"),
        (b"short text", "r.txt", "Almost no text"),
        (b"x" * (6 * 1024 * 1024), "r.pdf", "5 MB"),
        (b"old word file contents" * 20, "r.doc", ".docx"),
    ],
    # short names: pytest otherwise names a case after its data, and printing a
    # 6 MB name stalls the CI log
    ids=["empty", "wrong-type", "too-short", "too-big", "old-word-format"],
)
def test_unreadable_files_get_a_clear_reason(data, name, message):
    with pytest.raises(UnreadableResume, match=message):
        extract_text(data, name)


# ---------- profile ----------


def test_sections_and_dates():
    sections = split_sections(STRONG)
    assert {"education", "experience", "projects", "skills"} <= set(sections)
    ranges = date_ranges(sections["experience"], TODAY)
    assert ranges == [((2025, 5), (2025, 8)), ((2026, 1), (2026, 4))]
    assert months_of_experience(ranges) == 8
    # overlapping jobs count once
    assert months_of_experience([((2024, 1), (2024, 6)), ((2024, 4), (2024, 9))]) == 9


def test_profile_of_a_student():
    profile = build_profile(STRONG, TODAY)
    assert profile.level == "student"
    assert profile.target_levels[0] == "intern"
    assert profile.education["graduation"] == "2027-04" and profile.education["degree"] == "bachelors"
    assert "computer science" in profile.education["fields"]
    assert {"Python", "React", "TypeScript", "Spark", "Airflow", "Docker"} <= set(profile.skills)
    assert profile.contact == {
        "email": True,
        "phone": True,
        "linkedin": True,
        "github": True,
        "website": False,
    }
    assert profile.bullet_count == 6 and profile.quantified_bullets == 6
    assert profile.top_role in ("sde", "data_eng")


def test_profile_of_an_experienced_developer():
    text = (
        "Experience\nSenior Software Engineer | Shopify | Jan 2018 - Present\n• Led a team of 6 building payments\n"
        "Education\nBachelor of Science, University of Toronto, 2010 - 2014\nSkills\nRuby, Go, Kubernetes"
    )
    profile = build_profile(text, TODAY)
    assert profile.months_experience >= 100
    assert profile.level == "senior"


# ---------- checks ----------


def test_a_strong_resume_has_few_checks():
    profile = build_profile(STRONG, TODAY)
    checks, strength = check_resume(STRONG, profile)
    assert strength >= 85
    assert not [c for c in checks if c.severity == "high"]


def test_a_weak_resume_gets_specific_fixes():
    profile = build_profile(WEAK, TODAY)
    checks, strength = check_resume(WEAK, profile)
    ids = {c.id for c in checks}
    assert {"no_email", "weak_verbs", "few_skills", "buzzwords", "no_linkedin"} <= ids
    weak = next(c for c in checks if c.id == "weak_verbs")
    assert "Responsible for helping customers" in weak.examples
    assert strength < 50
    assert checks[0].severity == "high"  # most important first


# ---------- scoring ----------

JOB = {
    "id": 1,
    "title": "Data Engineering Intern, Payments",
    "level": "intern",
    "role_type": "data_eng",
    "skills": ["Python", "SQL", "Spark", "Airflow", "Kafka"],
    "signals": {"risk": "low", "days_open": 3},
}


def test_strong_match_lists_what_matches_and_whats_missing():
    profile = build_profile(STRONG, TODAY).to_dict()
    match = score_job(profile, JOB)
    assert match.score >= 75 and match.label in ("Strong fit", "Excellent fit")
    assert match.skills_matched == ["Python", "SQL", "Spark", "Airflow"]
    assert match.skills_missing == ["Kafka"]
    assert any("Kafka" in s for s in match.suggestions)
    assert any("payments" in s for s in match.suggestions)  # title word not in the resume


def test_level_and_role_mismatches_lower_the_score_and_say_why():
    profile = build_profile(STRONG, TODAY).to_dict()
    senior_hw = {
        **JOB,
        "title": "Senior ASIC Design Engineer",
        "level": "senior",
        "role_type": "hardware",
        "skills": ["Verilog", "SystemVerilog", "UVM"],
    }
    match = score_job(profile, senior_hw)
    assert match.score < 35 and match.label == "Long shot"
    assert any("senior role" in s for s in match.suggestions)
    assert any("hardware and embedded" in s for s in match.suggestions)


def test_jobs_without_descriptions_are_scored_cautiously():
    profile = build_profile(STRONG, TODAY).to_dict()
    match = score_job(profile, {**JOB, "skills": []})
    assert match.limited and match.score <= 75 and match.notes


def test_high_risk_jobs_say_so():
    profile = build_profile(STRONG, TODAY).to_dict()
    match = score_job(profile, {**JOB, "signals": {"risk": "high"}})
    assert any("ghost risk" in s for s in match.suggestions)


def test_recommendations_rank_by_fit():
    profile = build_profile(STRONG, TODAY).to_dict()
    jobs = [
        {**JOB, "id": 1},
        {**JOB, "id": 2, "title": "Senior Staff Engineer", "level": "senior", "skills": ["Go"]},
        {
            **JOB,
            "id": 3,
            "title": "Software Engineer Intern",
            "role_type": "sde",
            "skills": ["React", "TypeScript"],
        },
    ]
    ranked = [job["id"] for job, _ in recommend(profile, jobs, limit=3)]
    assert ranked[-1] == 2


def test_title_words_drop_generic_terms():
    assert title_words("Senior Software Engineer, Backend Payments (Toronto)") == ["backend", "payments"]


# ---------- builder documents ----------


def test_builder_documents_become_analyzable_text():
    doc = ResumeDocument.model_validate(
        {
            "contact": {
                "name": "Jordan Lee",
                "email": "j@example.com",
                "links": [{"label": "GitHub", "url": "github.com/jl"}],
            },
            "education": [{"school": "Waterloo", "degree": "BCS", "end": "Expected April 2027"}],
            "experience": [
                {
                    "title": "Software Developer Co-op",
                    "company": "Maple Bank",
                    "start": "May 2025",
                    "end": "Aug 2025",
                    "bullets": ["Built a React dashboard used by 40 analysts"],
                }
            ],
            "skills": [{"label": "Languages", "items": "Python, SQL"}],
        }
    )
    text = document_to_text(doc)
    assert "Experience" in text and "• Built a React dashboard" in text
    profile = build_profile(text, TODAY)
    assert {"Python", "SQL", "React"} <= set(profile.skills)
    assert profile.months_experience == 4


@pytest.mark.skipif(not os.environ.get("RUN_AI_TESTS"), reason="calls the real AI API")
def test_ai_feedback_live():  # pragma: no cover
    from tracker.resume import ai

    result = ai.feedback(STRONG, "Data Engineering Intern", "Acme", "Python, Spark, Kafka, Airflow pipelines")
    assert result.summary
